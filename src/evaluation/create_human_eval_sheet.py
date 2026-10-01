import json
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, Alignment
from openpyxl.utils import get_column_letter


# ============================================================
# Configuration
# ============================================================

INPUT = Path("results/benchmark/human_eval_set.json")
OUTPUT = Path("results/benchmark/human_evaluation.xlsx")


# ============================================================
# Load evaluation set
# ============================================================

with open(INPUT, "r", encoding="utf-8") as f:
    data = json.load(f)


# ============================================================
# Create workbook
# ============================================================

wb = Workbook()
ws = wb.active

# This workbook is a manual human-review template.
# The reported 30-example scores in the project are from
# AI-assisted LLM-as-judge evaluation, not this worksheet.
ws.title = "Human Review Template"


# ============================================================
# Headers
# ============================================================

headers = [
    "Eval ID",
    "Customer Query",
    "Base Response",
    "NexaTune Response",
    "Base Correctness (1-5)",
    "NexaTune Correctness (1-5)",
    "Base Relevance (1-5)",
    "NexaTune Relevance (1-5)",
    "Base Helpfulness (1-5)",
    "NexaTune Helpfulness (1-5)",
    "Base Professionalism (1-5)",
    "NexaTune Professionalism (1-5)",
    "Base Factuality (1-5)",
    "NexaTune Factuality (1-5)",
    "Overall Preference",
    "Notes",
]

ws.append(headers)


# ============================================================
# Populate evaluation examples
# ============================================================

for example in data["examples"]:
    ws.append([
        example["eval_id"],
        example["instruction"],
        example["base_prediction"],
        example["finetuned_prediction"],
        "",
        "",
        "",
        "",
        "",
        "",
        "",
        "",
        "",
        "",
        "",
        "",
    ])


# ============================================================
# Header formatting
# ============================================================

for cell in ws[1]:
    cell.font = Font(bold=True)
    cell.alignment = Alignment(
        horizontal="center",
        vertical="center",
        wrap_text=True,
    )


# ============================================================
# Wrap long text
# ============================================================

for row in ws.iter_rows(min_row=2):
    for cell in row:
        cell.alignment = Alignment(
            vertical="top",
            wrap_text=True,
        )


# ============================================================
# Practical column widths
# ============================================================

widths = {
    1: 10,
    2: 45,
    3: 65,
    4: 65,
    5: 22,
    6: 24,
    7: 20,
    8: 22,
    9: 22,
    10: 24,
    11: 24,
    12: 26,
    13: 20,
    14: 22,
    15: 20,
    16: 45,
}

for column, width in widths.items():
    ws.column_dimensions[get_column_letter(column)].width = width


# ============================================================
# Worksheet usability
# ============================================================

ws.freeze_panes = "A2"
ws.auto_filter.ref = ws.dimensions

# Increase row height for readability.
for row in range(2, ws.max_row + 1):
    ws.row_dimensions[row].height = 110


# ============================================================
# Save workbook
# ============================================================

OUTPUT.parent.mkdir(parents=True, exist_ok=True)
wb.save(OUTPUT)


# ============================================================
# Summary
# ============================================================

print("Human review template created.")
print(f"Examples: {len(data['examples'])}")
print(f"Saved to: {OUTPUT}")
print("Note: This workbook is a manual human-review template.")
print("Reported evaluation scores are from AI-assisted LLM-as-judge evaluation.")